from __future__ import annotations

import argparse
import html
from pathlib import Path

from pptx import Presentation

RISK_TERMS = {
    "old_version": ["拟研究", "计划采用", "预期成果", "项目进度", "研究计划"],
    "method_mismatch": ["ORR", "RDE", "半波电位"],
    "overclaim": ["最终机制", "绝对最优", "完全证明"],
}


def extract_slide_text(pptx_path: Path) -> list[dict]:
    prs = Presentation(pptx_path)
    slides = []
    for index, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = "\n".join(
                    p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()
                )
                if text:
                    texts.append(text)
        slides.append({"page": index, "text": "\n\n".join(texts)})
    return slides


def classify_risks(text: str) -> list[tuple[str, str]]:
    findings = []
    for category, terms in RISK_TERMS.items():
        for term in terms:
            if term in text:
                findings.append((category, term))
    return findings


def build_html_report(slides: list[dict], output: Path) -> None:
    rows = []
    for slide in slides:
        risks = classify_risks(slide["text"])
        priority = "P0" if any(c in {"old_version", "method_mismatch"} for c, _ in risks) else ("P1" if risks else "P2")
        risk_text = ", ".join(f"{c}:{t}" for c, t in risks) or "No obvious risk term"
        rows.append(
            f"<tr><td>{slide['page']}</td><td>{priority}</td>"
            f"<td>{html.escape(risk_text)}</td>"
            f"<td><pre>{html.escape(slide['text'][:1200])}</pre></td></tr>"
        )
    doc = f"""<!doctype html><html><head><meta charset='utf-8'><title>PPT Review Report</title>
<style>body{{font-family:Arial,sans-serif;background:#f6f8fb;padding:24px}}table{{width:100%;border-collapse:collapse;background:white}}td,th{{border:1px solid #ddd;padding:8px;vertical-align:top}}pre{{white-space:pre-wrap}}</style></head><body>
<h1>PPT Review Report</h1><table><tr><th>Page</th><th>Priority</th><th>Risk</th><th>Text</th></tr>{''.join(rows)}</table></body></html>"""
    output.write_text(doc, encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate a simple PPT risk review report.")
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--out", type=Path, default=Path("review_report.html"))
    args = parser.parse_args()
    slides = extract_slide_text(args.pptx)
    build_html_report(slides, args.out)
    print(f"Report written to {args.out}")


if __name__ == "__main__":
    main()
